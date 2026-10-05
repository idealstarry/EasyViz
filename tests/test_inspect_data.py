"""Directory intake preserves sources and makes only bounded candidate claims."""
import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch


SCRIPT = Path(__file__).resolve().parents[1] / "skills/easyviz/scripts/inspect_data.py"
loader = importlib.util.spec_from_file_location("intake", SCRIPT)
intake = importlib.util.module_from_spec(loader)
loader.loader.exec_module(intake)


class InspectDataTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="easyviz-intake-")
        self.root = Path(self.temp.name)
        self.data = self.root / "data"
        self.data.mkdir()
        self.source = self.data / "observations.csv"
        self.source.write_text("sample_id,condition,measurement,other\n001,A,1.0,9\n002,A,2.0,8\n003,B,NA,7\n004,B,4.0,6\n", encoding="utf-8")
        self.out = self.root / "exploration"

    def tearDown(self):
        self.temp.cleanup()

    def cli(self, *extra, input_path=None, out=None):
        return subprocess.run([sys.executable, str(SCRIPT), "--input", str(input_path or self.data),
                               "--out", str(out or self.out), *map(str, extra)],
                              cwd=self.root, capture_output=True, text=True)

    def test_mixed_directory_inventory_preserves_identifiers_and_sources(self):
        sub = self.data / "nested"
        sub.mkdir()
        (sub / "metadata.tsv").write_text("sample_id\tbatch\n001\tfirst\n002\tsecond\n")
        (self.data / "README.md").write_text("not a table")
        (self.data / "broken.csv").write_text("a,b\n1,2,3\n")
        before = {path: path.read_bytes() for path in self.data.rglob("*") if path.is_file()}
        result = self.cli()
        self.assertEqual(result.returncode, 0, result.stderr)
        manifest = json.loads((self.out / "manifest.json").read_text())
        self.assertEqual(len(manifest["tables"]), 2)
        self.assertEqual(len(manifest["errors"]), 1)
        table = next(table for table in manifest["tables"] if table["path"] == self.source.name)
        self.assertEqual(table["sample_rows"][0]["sample_id"], "001")
        self.assertEqual(table["columns"][0]["type_hint"], "identifier_candidate")
        self.assertEqual(table["columns"][0]["storage"], "string")
        measurement = next(column for column in table["columns"] if column["name"] == "measurement")
        self.assertEqual(measurement["missing"], 1)
        self.assertEqual(measurement["descriptive"]["n"], 3)
        for path, content in before.items():
            self.assertEqual(path.read_bytes(), content)
        self.assertFalse(manifest["inference_run"])
        self.assertFalse(manifest["sources_modified"])
        self.assertTrue((self.out / "exploration.md").is_file())

    def test_candidates_do_not_invent_design_or_run_significance(self):
        result = self.cli()
        self.assertEqual(result.returncode, 0, result.stderr)
        options = json.loads((self.out / "analysis-options.json").read_text())
        self.assertEqual(options["track"], "create")
        self.assertFalse(options["inference_run"])
        table = options["tables"][0]
        self.assertIsNone(table["declared_design"])
        distribution = next(option for option in table["options"] if option["id"] == "distribution")
        plan = distribution["descriptive_analysis_plan"]
        self.assertEqual(plan["comparisons"][0]["method"], "descriptive")
        self.assertIsNone(plan["design"]["unit"])
        self.assertFalse(plan["design"]["confirmed"])
        for option in table["options"]:
            self.assertTrue(option["inference"]["questions"])
            self.assertNotIn("p_value", option)
        self.assertLessEqual(len(table["options"]), 3)

    def test_numeric_fields_without_adopted_roles_do_not_select_joint_chart_or_correlation(self):
        self.source.write_text("sample_id,weight,length,height\n001,2,10,20\n002,3,11,22\n003,4,12,24\n")
        before = self.source.read_bytes()
        _, recommendations = intake.inspect(self.source, self.out)
        options = recommendations["tables"][0]["options"]
        hint = next(option for option in options if option["id"] == "numeric_pair_eligibility")
        self.assertEqual(hint["recommended_charts"], [])
        self.assertNotIn("proposed_fields", hint)
        self.assertNotIn("methods_to_consider_after_design", hint["inference"])
        self.assertEqual(hint["inference"]["status"], "blocked")
        self.assertNotIn("association", [option["id"] for option in options])
        self.assertNotIn("scatter", [chart for option in options for chart in option["recommended_charts"]])
        self.assertNotIn("pearson", json.dumps(options))
        self.assertNotIn("spearman", json.dumps(options))
        self.assertEqual(self.source.read_bytes(), before)

    def test_question_text_and_partial_x_role_do_not_fill_in_y_or_choose_joint_chart(self):
        path = self.root / "design.json"
        path.write_text(json.dumps({"table": self.source.name, "fields": {"x": "other"},
                                   "question": "Plot the before/after correlation using a scatter plot"}))
        _, recommendations = intake.inspect(self.source, self.out, path)
        options = recommendations["tables"][0]["options"]
        hint = next(option for option in options if option["id"] == "numeric_pair_eligibility")
        self.assertNotIn("proposed_fields", hint)
        self.assertEqual(hint["recommended_charts"], [])
        self.assertNotIn("association", [option["id"] for option in options])
        self.assertNotIn("pearson", json.dumps(options))
        self.assertNotIn("spearman", json.dumps(options))

    def test_explicit_xy_mapping_is_preserved_without_claiming_reading_purpose_adoption(self):
        path = self.root / "design.json"
        path.write_text(json.dumps({"table": self.source.name, "fields": {"x": "other", "y": "measurement"},
                                   "question": "Explore this table"}))
        _, recommendations = intake.inspect(self.source, self.out, path)
        options = recommendations["tables"][0]["options"]
        association = next(option for option in options if option["id"] == "association")
        self.assertEqual(association["recommended_charts"], ["scatter"])
        self.assertEqual(association["proposed_fields"], {"x": "other", "y": "measurement"})
        self.assertEqual(association["field_candidates"]["x"], ["other"])
        self.assertEqual(association["field_candidates"]["y"], ["measurement"])
        self.assertEqual(association["mapping_status"], "declared")
        self.assertIn("requires an adopted reading purpose", association["purpose_status"])
        self.assertEqual(association["inference"]["status"], "not_selected")
        self.assertNotIn("numeric_pair_eligibility", [option["id"] for option in options])

    def test_confirmed_paired_wide_candidates_require_schema_and_traceable_preparation(self):
        self.source.write_text("sample_id,cohort,before,after,age\n001,A,1.25,1.75,20\n002,A,2.5,2.0,30\n003,B,3.0,NA,40\n")
        before = self.source.read_bytes()
        path = self.root / "design.json"
        declaration = {"table": self.source.name, "row_kind": "observations",
                       "design": {"unit": "sample_id", "unit_definition": "one participant with measurements in multiple conditions", "structure": "paired", "confirmed": True}}
        path.write_text(json.dumps(declaration))
        manifest, recommendations = intake.inspect(self.source, self.out, path)
        options = recommendations["tables"][0]["options"]
        identities = [option["id"] for option in options]
        self.assertLess(identities.index("paired_wide_candidate"), identities.index("numeric_pair_eligibility"))
        self.assertNotIn("paired_candidate", identities, "Cohort labels must not become paired conditions")
        candidate = next(option for option in options if option["id"] == "paired_wide_candidate")
        self.assertEqual(candidate["recommended_charts"], ["paired"])
        self.assertEqual(candidate["field_candidates"]["unit"], ["sample_id"])
        self.assertNotIn("proposed_fields", candidate)
        self.assertNotIn("descriptive_analysis_plan", candidate)
        preparation = candidate["preparation"]
        self.assertEqual(preparation["status"], "required_before_plotting")
        self.assertEqual(preparation["unit_verification"]["status"], "one_row_per_unit_in_complete_source")
        self.assertEqual(preparation["unit_verification"]["unique_nonmissing"], 3)
        self.assertIn("unrelated numeric covariates", preparation["required_steps"][0])
        self.assertIn("missing/unmatched", preparation["required_steps"][1])
        self.assertIn("source row/column or cell", preparation["required_steps"][2])
        self.assertIn("unchanged measurements", preparation["required_steps"][2])
        self.assertEqual(candidate["inference"]["status"], "blocked")
        self.assertNotIn("methods_to_consider_after_design", candidate["inference"])
        self.assertEqual(manifest["tables"][0]["sample_rows"][0]["sample_id"], "001")
        self.assertEqual(manifest["tables"][0]["sample_rows"][2]["after"], "NA")
        self.assertEqual(self.source.read_bytes(), before)
        self.assertEqual({file.name for file in self.out.iterdir()}, {"manifest.json", "analysis-options.json", "exploration.md"})

    def test_paired_wide_prefix_check_does_not_claim_full_source_unit_verification(self):
        self.source.write_text("sample_id,before,after\n001,1,2\n002,2,3\n001,3,4\n")
        path = self.root / "design.json"
        path.write_text(json.dumps({"table": self.source.name,
                                   "design": {"unit": "sample_id", "unit_definition": "one participant", "structure": "paired", "confirmed": True}}))
        with patch.dict(intake.LIMITS, rows_per_table=2):
            _, recommendations = intake.inspect(self.source, self.out, path)
        candidate = next(option for option in recommendations["tables"][0]["options"] if option["id"] == "paired_wide_candidate")
        self.assertIn("inspected_prefix_only", candidate["preparation"]["unit_verification"]["status"])
        self.assertEqual(candidate["preparation"]["status"], "required_before_plotting")

    def test_wide_pairing_is_not_inferred_from_names_or_dirty_unit_keys(self):
        path = self.root / "design.json"
        examples = [("001,1,2\n002,2,3\n", "paired", False),
                    ("001,1,2\n002,2,3\n", "independent", True),
                    ("001,1,2\n001,2,3\n", "paired", True),
                    ("001,1,2\n,2,3\n", "paired", True)]
        for index, (rows, structure, confirmed) in enumerate(examples):
            with self.subTest(structure=structure, confirmed=confirmed, rows=rows):
                self.source.write_text("sample_id,before,after\n" + rows)
                before = self.source.read_bytes()
                path.write_text(json.dumps({"table": self.source.name,
                                           "design": {"unit": "sample_id", "unit_definition": "one participant", "structure": structure, "confirmed": confirmed}}))
                _, recommendations = intake.inspect(self.source, self.root / f"check-{index}", path)
                self.assertNotIn("paired_wide_candidate", [option["id"] for option in recommendations["tables"][0]["options"]])
                self.assertEqual(self.source.read_bytes(), before)

    def test_inventory_only_preserves_source_profiles_without_track_or_recommendations(self):
        original = self.source.read_bytes()
        result = self.cli("--inventory-only")
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual({path.name for path in self.out.iterdir()}, {"manifest.json", "exploration.md"})
        manifest = json.loads((self.out / "manifest.json").read_text())
        self.assertIsNone(manifest["track"])
        self.assertEqual(manifest["mode"], "inventory")
        self.assertEqual(manifest["tables"][0]["sample_rows"][0]["sample_id"], "001")
        self.assertEqual(manifest["tables"][0]["rows_inspected"], 4)
        self.assertFalse(manifest["inference_run"])
        report = (self.out / "exploration.md").read_text()
        self.assertIn("active plotting track remains unchanged", report)
        self.assertNotIn("analysis-options.json", report)
        self.assertNotIn("candidate reading tasks", report)
        self.assertEqual(self.source.read_bytes(), original)
        # The shared inventory path must not even invoke the create recommender.
        with patch.object(intake, "_options", side_effect=AssertionError("create recommender invoked")):
            _, recommendations = intake.inspect(self.source, self.root / "api-inventory", inventory_only=True)
        self.assertIsNone(recommendations)

    def test_safe_scan_excludes_symlinks_hidden_output_and_dependencies(self):
        outside = self.root / "outside.csv"
        outside.write_text("secret\nexternal\n")
        (self.data / "escape.csv").symlink_to(outside)
        for name in (".hidden", "outputs", "venv", "node_modules"):
            directory = self.data / name
            directory.mkdir()
            (directory / "ignored.csv").write_text("x\n1\n")
        result = self.cli(out=self.data / "review-run")
        self.assertEqual(result.returncode, 0, result.stderr)
        manifest = json.loads((self.data / "review-run/manifest.json").read_text())
        self.assertEqual([table["path"] for table in manifest["tables"]], ["observations.csv"])
        self.assertTrue(any(entry["path"] == "escape.csv" and "symlink" in entry["reason"] for entry in manifest["skipped"]))

    def test_outputs_never_replace_sources_or_prior_review(self):
        for out in (self.data, self.source):
            result = self.cli(out=out)
            self.assertEqual(result.returncode, 2)
            self.assertNotIn("Traceback", result.stderr)
        self.out.mkdir()
        sentinel = self.out / "manifest.json"
        sentinel.write_text("original")
        result = self.cli()
        self.assertEqual(result.returncode, 2)
        self.assertEqual(sentinel.read_text(), "original")

    def test_explicit_input_symlink_is_rejected_without_reading(self):
        link = self.root / "linked.csv"
        link.symlink_to(self.source)
        result = self.cli(input_path=link)
        self.assertEqual(result.returncode, 2)
        self.assertIn("symlink", result.stderr)
        self.assertFalse(self.out.exists())

    def test_no_usable_table_has_actionable_error_and_no_partial_report(self):
        self.source.write_text("a,a\n1,2\n")
        result = self.cli()
        self.assertEqual(result.returncode, 2)
        self.assertIn("duplicate column names", result.stderr)
        self.assertNotIn("Traceback", result.stderr)
        self.assertFalse(self.out.exists())

    def test_empty_and_non_utf8_tables_are_reported_in_mixed_input(self):
        (self.data / "empty.csv").write_text("")
        (self.data / "legacy.tsv").write_bytes(b"id\n\xff\n")
        result = self.cli()
        self.assertEqual(result.returncode, 0, result.stderr)
        errors = json.loads((self.out / "manifest.json").read_text())["errors"]
        self.assertEqual(len(errors), 2)
        self.assertTrue(any("UTF-8" in error["reason"] for error in errors))

    def test_declared_roles_are_preserved_without_selecting_inference(self):
        declaration = {"table": "observations.csv", "question": "Compare the adopted measurement",
                       "fields": {"group": "condition", "value": "measurement"},
                       "design": {"unit": "sample_id", "unit_definition": "one independently sampled specimen", "structure": "independent", "confirmed": True},
                       "missing_tokens": ["", "NA"], "missing_policy": "complete_case"}
        design_path = self.root / "design.json"
        design_path.write_text(json.dumps(declaration))
        result = self.cli("--design", design_path)
        self.assertEqual(result.returncode, 0, result.stderr)
        table = json.loads((self.out / "analysis-options.json").read_text())["tables"][0]
        self.assertEqual(table["declared_design"], declaration)
        plan = table["options"][0]["descriptive_analysis_plan"]
        self.assertEqual(plan["design"], declaration["design"])
        self.assertEqual(plan["comparisons"][0]["fields"], declaration["fields"])
        self.assertEqual(plan["comparisons"][0]["method"], "descriptive")
        self.assertNotIn("paired_candidate", [option["id"] for option in table["options"]])
        methods = table["options"][0]["inference"]["methods_to_consider_after_design"]
        self.assertEqual(methods, ["welch or mannwhitney for independent two-group questions"])

    def test_missing_declared_column_fails_before_creating_output(self):
        path = self.root / "design.json"
        path.write_text(json.dumps({"table": "observations.csv", "fields": {"value": "imagined"}}))
        result = self.cli("--design", path)
        self.assertEqual(result.returncode, 2)
        self.assertIn("Declared columns not found", result.stderr)
        self.assertFalse(self.out.exists())

    def test_row_limits_make_profile_scope_explicit(self):
        with patch.dict(intake.LIMITS, rows_per_table=2):
            manifest, options = intake.inspect(self.source, self.out)
        table = manifest["tables"][0]
        self.assertFalse(table["rows_complete"])
        self.assertIsNone(table["row_count"])
        self.assertEqual(table["rows_inspected"], 2)
        self.assertIn("prefix", options["tables"][0]["profile_scope"])

    def test_duplicate_rows_are_reported_not_removed(self):
        self.source.write_text("sample_id,measurement\n001,2\n001,2\n002,3\n")
        manifest, _ = intake.inspect(self.source, self.out)
        self.assertEqual(manifest["tables"][0]["rows_inspected"], 3)
        self.assertEqual(manifest["tables"][0]["duplicate_rows_in_inspected_scope"], 1)

    def test_identifiers_and_nonnumeric_tables_do_not_become_measurements(self):
        self.source.write_text("sample_id,label\n001,case\n002,control\n")
        _, options = intake.inspect(self.source, self.out)
        self.assertEqual(options["tables"][0]["candidate_columns"]["numeric"], [])
        self.assertEqual(options["tables"][0]["options"][0]["id"], "table_structure")

    def test_messy_multitable_intake_exposes_questions_without_joining_or_inventing_roles(self):
        self.source.write_text("sample_id,condition,v1,v2\n001,A,1,5\n001,B,2,4\n001,B,3,3\n002,A,4,2\n002,B,5,1\n")
        metadata = self.data / "metadata.csv"
        metadata.write_text("sample_id,batch\n001,first\n002,second\n")
        summary = self.data / "summary.csv"
        summary.write_text("time (h),mean (mm),sd (mm),n\n0,2,0.3,10\n1,3,0.4,10\n")
        before = {path: path.read_bytes() for path in (self.source, metadata, summary)}
        manifest, recommendations = intake.inspect(self.data, self.out)
        tables = {table["table_id"]: table for table in manifest["tables"]}
        choices = {table["table_id"]: table for table in recommendations["tables"]}
        self.assertEqual(tables[self.source.name]["intake_signals"]["ambiguous_headers"], ["v1", "v2"])
        unit = tables[self.source.name]["intake_signals"]["identifier_candidates"][0]
        self.assertEqual(unit["repeated_identifiers"], 2)
        self.assertEqual(unit["rows_per_identifier_max"], 3)
        self.assertEqual(unit["group_patterns"], [{"group_column": "condition", "identifiers_in_multiple_levels": 2,
                                                 "repeated_identifier_level_cells": 1}])
        raw = choices[self.source.name]
        self.assertIn("paired_candidate", [option["id"] for option in raw["options"]])
        self.assertLess([option["id"] for option in raw["options"]].index("paired_candidate"),
                        [option["id"] for option in raw["options"]].index("numeric_pair_eligibility"))
        self.assertNotIn("association", [option["id"] for option in raw["options"]])
        self.assertEqual([item["id"] for item in raw["intake"]["priority_questions"]], ["repeated_unit", "field_meaning"])
        self.assertIn("technical repeats", raw["intake"]["priority_questions"][0]["question"])
        self.assertIn("ambiguous_headers", [item["id"] for item in raw["intake"]["cautions"]])
        self.assertEqual(choices[metadata.name]["options"][0]["id"], "table_structure")
        self.assertEqual(choices[summary.name]["options"][0]["id"], "supplied_summary")
        self.assertIn("timecourse", choices[summary.name]["options"][0]["recommended_charts"])
        self.assertNotIn("joins", manifest)
        self.assertEqual(tables[self.source.name]["sample_rows"][0]["sample_id"], "001")
        for path, content in before.items():
            self.assertEqual(path.read_bytes(), content)
        for table in recommendations["tables"]:
            self.assertLessEqual(len(table["intake"]["priority_questions"]), 3)
            self.assertFalse(any("pvalue" in option for option in table["options"]))

    def test_supplied_sd_and_n_do_not_become_raw_distribution_or_inference_plan(self):
        self.source.write_text("condition,mean,sd,n\nA,12,2,30\nB,14,3,40\n")
        manifest, recommendations = intake.inspect(self.source, self.out)
        table = recommendations["tables"][0]
        self.assertEqual(manifest["tables"][0]["intake_signals"]["summary_data"]["status"], "suspected")
        self.assertEqual(table["intake"]["status"], "suspected_summary_needs_context")
        self.assertEqual([option["id"] for option in table["options"]], ["supplied_summary"])
        option = table["options"][0]
        self.assertEqual(option["field_candidates"]["estimate"], ["mean"])
        self.assertEqual(option["field_candidates"]["supplied_n"], ["n"])
        self.assertEqual(option["inference"]["status"], "blocked")
        self.assertNotIn("descriptive_analysis_plan", option)
        self.assertNotIn("methods_to_consider_after_design", option["inference"])
        report = (self.out / "exploration.md").read_text()
        self.assertIn("do not test summary rows as raw replicates", report)
        self.assertIn("SD, SEM, CI", report)

    def test_bounds_route_to_interval_without_assuming_confidence_level(self):
        self.source.write_text("term,estimate,ci_lower,ci_upper\na,0.4,0.1,0.8\nb,-0.3,-0.5,0.1\n")
        _, recommendations = intake.inspect(self.source, self.out)
        option = recommendations["tables"][0]["options"][0]
        self.assertEqual(option["recommended_charts"], ["interval"])
        self.assertEqual(option["field_candidates"]["uncertainty"]["lower_bound"], ["ci_lower"])
        self.assertNotIn("95", json.dumps(option))
        self.assertNotIn("descriptive_analysis_plan", option)

    def test_sem_is_not_advertised_as_timecourse_sd(self):
        self.source.write_text("time,value,sem\n0,2,0.1\n1,3,0.2\n")
        _, recommendations = intake.inspect(self.source, self.out)
        option = recommendations["tables"][0]["options"][0]
        self.assertEqual(option["id"], "supplied_summary")
        self.assertNotIn("timecourse", option["recommended_charts"])
        self.assertEqual(option["field_candidates"]["uncertainty"]["standard_error"], ["sem"])
        self.assertEqual(option["field_candidates"]["estimate"], ["value"])

    def test_partial_bounds_do_not_turn_into_raw_observations_or_invent_missing_bounds(self):
        self.source.write_text("term,estimate,lower\na,0.4,0.1\nb,0.5,0.2\n")
        _, recommendations = intake.inspect(self.source, self.out)
        option = recommendations["tables"][0]["options"][0]
        self.assertEqual(option["id"], "supplied_summary")
        self.assertNotIn("interval", option["recommended_charts"])
        self.assertEqual(option["field_candidates"]["uncertainty"]["upper_bound"], [])
        self.assertNotIn("descriptive_analysis_plan", option)

    def test_bounds_without_estimates_request_roles_before_advertising_recipe(self):
        self.source.write_text("term,lower,upper\na,0.1,0.8\nb,-0.5,0.1\n")
        _, recommendations = intake.inspect(self.source, self.out)
        option = recommendations["tables"][0]["options"][0]
        self.assertEqual(option["id"], "supplied_summary")
        self.assertEqual(option["field_candidates"]["estimate"], [])
        self.assertEqual(option["recommended_charts"], [])
        self.assertNotIn("descriptive_analysis_plan", option)

    def test_name_hints_require_distinct_support_and_do_not_invent_unit_labels(self):
        self.source.write_text("sample_id,mean_cell_size\n001,2.0\n002,3.0\n")
        manifest, recommendations = intake.inspect(self.source, self.out)
        self.assertEqual(manifest["tables"][0]["intake_signals"]["summary_data"]["status"], "not_established")
        self.assertEqual(recommendations["tables"][0]["options"][0]["id"], "distribution")
        self.assertIsNone(manifest["tables"][0]["columns"][1]["unit_label_hint"])

    def test_unique_identifiers_do_not_suggest_pairing(self):
        _, recommendations = intake.inspect(self.source, self.out)
        self.assertNotIn("paired_candidate", [option["id"] for option in recommendations["tables"][0]["options"]])

    def test_explicit_observation_grain_resolves_summary_like_names(self):
        self.source.write_text("sample_id,condition,mean,sd\n001,A,2,0.1\n002,B,3,0.2\n")
        path = self.root / "design.json"
        declaration = {"table": self.source.name, "row_kind": "observations", "fields": {"value": "mean", "group": "condition"},
                       "design": {"unit": "sample_id", "unit_definition": "one independently sampled specimen whose measurement is an image mean", "structure": "independent", "confirmed": True}}
        path.write_text(json.dumps(declaration))
        manifest, recommendations = intake.inspect(self.source, self.out, path)
        self.assertEqual(manifest["tables"][0]["intake_signals"]["summary_data"]["status"], "suspected")
        table = recommendations["tables"][0]
        self.assertEqual(table["intake"]["row_kind"], "observations")
        self.assertEqual(table["options"][0]["id"], "distribution")
        self.assertEqual(table["options"][0]["descriptive_analysis_plan"]["comparisons"][0]["fields"], declaration["fields"])
        self.assertEqual(table["declared_design"]["row_kind"], "observations")

    def test_explicit_summary_grain_blocks_raw_plan_even_without_named_summary_columns(self):
        self.source.write_text("category,v1,v2,v3\nA,2,1,3\nB,4,3,5\n")
        path = self.root / "design.json"
        path.write_text(json.dumps({"table": self.source.name, "row_kind": "summaries", "fields": {"value": "v1", "group": "category"}}))
        _, recommendations = intake.inspect(self.source, self.out, path)
        table = recommendations["tables"][0]
        self.assertEqual(table["intake"]["status"], "declared_summaries")
        self.assertEqual(table["options"][0]["field_candidates"]["estimate"], ["v1"])
        self.assertNotIn("descriptive_analysis_plan", table["options"][0])

    def test_inventory_only_has_summary_diagnostics_but_no_create_guidance(self):
        self.source.write_text("sample_id,mean,sd,n\n001,2,0.2,8\n001,3,0.3,9\n")
        with patch.object(intake, "_intake_guidance", side_effect=AssertionError("guidance invoked")):
            manifest, recommendations = intake.inspect(self.source, self.out, inventory_only=True)
        self.assertIsNone(recommendations)
        self.assertEqual(manifest["tables"][0]["intake_signals"]["summary_data"]["status"], "suspected")
        serialized = json.dumps(manifest)
        for forbidden in ("recommended_charts", "priority_questions", "descriptive_analysis_plan", "methods_to_consider"):
            self.assertNotIn(forbidden, serialized)
        report = (self.out / "exploration.md").read_text()
        self.assertNotIn("Resolve these points", report)
        self.assertNotIn("interval", report)

    def test_invalid_row_kind_fails_without_partial_output(self):
        path = self.root / "design.json"
        for invalid in ("independent", [], {}, None, True):
            with self.subTest(row_kind=invalid):
                path.write_text(json.dumps({"table": self.source.name, "row_kind": invalid}))
                result = self.cli("--design", path)
                self.assertEqual(result.returncode, 2)
                self.assertIn("row_kind", result.stderr)
                self.assertNotIn("Traceback", result.stderr)
                self.assertFalse(self.out.exists())

    def test_unit_labels_are_recorded_literally_without_adopting_scientific_meaning(self):
        self.source.write_text("sample_id,signal [AU],value (mm)\n001,2,3\n002,4,5\n")
        manifest, _ = intake.inspect(self.source, self.out)
        columns = manifest["tables"][0]["columns"]
        self.assertEqual(columns[1]["unit_label_hint"], "AU")
        self.assertEqual(columns[2]["unit_label_hint"], "mm")
        self.assertEqual(columns[1]["type_hint"], "numeric")
        self.assertEqual(columns[1]["storage"], "string")
        self.assertEqual(manifest["tables"][0]["grain"]["scientific_unit"], "unconfirmed")

    def test_attached_unit_labels_do_not_hide_ambiguous_roles_or_coordinate_hints(self):
        self.source.write_text("time(h),value[mm],sd[mm]\n0,2,0.1\n1,3,0.2\n")
        manifest, recommendations = intake.inspect(self.source, self.out)
        self.assertEqual(manifest["tables"][0]["intake_signals"]["ambiguous_headers"], ["value[mm]"])
        option = recommendations["tables"][0]["options"][0]
        self.assertEqual(option["field_candidates"]["estimate"], ["value[mm]"])
        self.assertEqual(option["field_candidates"]["x"], ["time(h)"])
        self.assertIn("timecourse", option["recommended_charts"])

    def test_declared_independent_unit_repetitions_remain_visible(self):
        self.source.write_text("sample_id,condition,measurement\n001,A,1\n001,A,2\n002,B,3\n")
        path = self.root / "design.json"
        path.write_text(json.dumps({"table": self.source.name, "fields": {"value": "measurement", "group": "condition"},
                                   "design": {"unit": "sample_id", "unit_definition": "one specimen", "structure": "independent", "confirmed": True}}))
        _, recommendations = intake.inspect(self.source, self.out, path)
        table = recommendations["tables"][0]
        self.assertIn("declared_unit_repeated", [item["id"] for item in table["intake"]["cautions"]])
        self.assertNotIn("paired_candidate", [option["id"] for option in table["options"]])

    def test_numeric_declared_unit_is_not_an_automatic_measurement_and_plan_executes(self):
        self.source.write_text("well,condition,measurement,other\n1,A,10,5\n2,A,12,6\n3,B,14,7\n4,B,16,8\n")
        declaration = {"table": self.source.name, "row_kind": "observations", "fields": {"group": "condition"},
                       "design": {"unit": "well", "unit_definition": "one independently sampled well", "structure": "independent", "confirmed": True}}
        design_path = self.root / "design.json"
        design_path.write_text(json.dumps(declaration))
        _, recommendations = intake.inspect(self.source, self.out, design_path)
        table = recommendations["tables"][0]
        self.assertEqual(table["candidate_columns"]["numeric"], ["measurement", "other"])
        self.assertIn("well", table["candidate_columns"]["identifier"])
        for option in table["options"]:
            self.assertNotIn("well", option.get("proposed_fields", {}).values())
            for role in ("value", "x", "y", "group"):
                self.assertNotIn("well", option["field_candidates"].get(role, []))
        plan = table["options"][0]["descriptive_analysis_plan"]
        plan_path = self.root / "adopted-plan.json"
        plan_path.write_text(json.dumps(plan))
        result = subprocess.run([sys.executable, str(SCRIPT.with_name("analyze.py")), "--data", str(self.source),
                                 "--plan", str(plan_path), "--out", str(self.root / "analysis")],
                                cwd=self.root, capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(plan["comparisons"][0]["fields"], {"value": "measurement", "group": "condition"})

    def test_categorical_declared_unit_is_not_an_automatic_group(self):
        self.source.write_text("well,condition,measurement\nW1,A,10\nW2,A,12\nW3,B,14\nW4,B,16\n")
        design_path = self.root / "design.json"
        design_path.write_text(json.dumps({"table": self.source.name, "row_kind": "observations",
                                          "design": {"unit": "well", "unit_definition": "one specimen", "structure": "independent", "confirmed": True}}))
        _, recommendations = intake.inspect(self.source, self.out, design_path)
        table = recommendations["tables"][0]
        self.assertEqual(table["candidate_columns"]["group"], ["condition"])
        self.assertEqual(table["options"][0]["proposed_fields"], {"value": "measurement", "group": "condition"})

    def test_declared_unit_cannot_also_be_declared_group_or_measurement(self):
        design_path = self.root / "design.json"
        for role in ("group", "value", "x", "y"):
            with self.subTest(role=role):
                design_path.write_text(json.dumps({"table": self.source.name, "fields": {role: "sample_id"},
                                                   "design": {"unit": "sample_id", "unit_definition": "one specimen", "structure": "independent", "confirmed": True}}))
                result = self.cli("--design", design_path)
                self.assertEqual(result.returncode, 2)
                self.assertIn("Unit IDs cannot also", result.stderr)
                self.assertFalse(self.out.exists())

    def test_custom_missing_tokens_keep_na_category(self):
        self.source.write_text("condition,value\nNA,1\nB,2\n")
        path = self.root / "design.json"
        path.write_text(json.dumps({"table": "observations.csv", "missing_tokens": [""]}))
        manifest, _ = intake.inspect(self.source, self.out, path)
        self.assertEqual(manifest["tables"][0]["columns"][0]["missing"], 0)
        self.assertIn("NA", manifest["tables"][0]["columns"][0]["levels"])

    def test_xlsx_sheets_formula_text_and_string_identifiers(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl is optional outside the installed runtime")
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Observations"
        sheet.append(["sample_id", "measurement", "formula"])
        sheet.append(["001", 2, "=1+1"])
        second = workbook.create_sheet("Metadata")
        second.append(["sample_id", "batch"])
        second.append(["001", "first"])
        path = self.data / "book.xlsx"
        workbook.save(path)
        result = self.cli(input_path=path)
        self.assertEqual(result.returncode, 0, result.stderr)
        manifest = json.loads((self.out / "manifest.json").read_text())
        self.assertEqual(len(manifest["tables"]), 2)
        observed = manifest["tables"][0]
        self.assertEqual(observed["sample_rows"][0]["sample_id"], "001")
        self.assertEqual(observed["sample_rows"][0]["formula"], "=1+1")
        self.assertEqual(observed["columns"][2]["formula_cells"], 1)

    def test_corrupt_xlsx_in_mixed_directory_is_recorded(self):
        (self.data / "bad.xlsx").write_bytes(b"not a zip workbook")
        result = self.cli()
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(len(json.loads((self.out / "manifest.json").read_text())["errors"]), 1)

    def test_valid_xlsx_sheet_survives_unusable_sheet_in_same_workbook(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl is optional outside the installed runtime")
        workbook = openpyxl.Workbook()
        workbook.active.title = "Bad"
        workbook.active.append(["duplicate", "duplicate"])
        workbook.active.append([1, 2])
        good = workbook.create_sheet("Good")
        good.append(["value"])
        good.append([3])
        path = self.data / "sheets.xlsx"
        workbook.save(path)
        manifest, _ = intake.inspect(path, self.out)
        self.assertEqual([table["sheet"] for table in manifest["tables"]], ["Good"])
        self.assertEqual(len(manifest["errors"]), 1)

    def test_missing_tokens_match_analysis_exactly_without_trimming(self):
        self.source.write_text("measurement\n NA \n2\n")
        manifest, _ = intake.inspect(self.source, self.out)
        column = manifest["tables"][0]["columns"][0]
        self.assertEqual(column["missing"], 0)
        self.assertEqual(column["nonmissing"], 2)
        self.assertEqual(column["numeric_parseable"], 1)

    def test_descriptive_candidate_executes_through_analysis_cli(self):
        script = SCRIPT.with_name("analyze.py")
        self.assertTrue(script.is_file(), "Directory intake requires the advertised analysis entry point")
        _, options = intake.inspect(self.source, self.out)
        plan = options["tables"][0]["options"][0]["descriptive_analysis_plan"]
        # Adopting complete_case is a recorded user/Agent choice, not a hidden
        # effect of the inventory step. Unknown units remain unknown.
        plan["missing_policy"] = "complete_case"
        path = self.root / "adopted-plan.json"
        path.write_text(json.dumps(plan))
        result = subprocess.run([sys.executable, str(script), "--data", str(self.source),
                                 "--plan", str(path), "--out", str(self.root / "analysis")],
                                cwd=self.root, capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        report = json.loads((self.root / "analysis/results.json").read_text())
        comparison = report["comparisons"][0]
        self.assertEqual(comparison["method"], "descriptive")
        self.assertIsNone(comparison["pvalue"])
        self.assertIsNone(comparison["counts"]["independent_unit_count"])
        self.assertEqual(comparison["counts"]["included_rows"], 3)
        self.assertIn("001", (self.root / "analysis/analyzed-data.csv").read_text())

    def test_extreme_finite_numbers_do_not_emit_nonfinite_json(self):
        self.source.write_text("measurement\n1e308\n1e308\n")
        manifest, _ = intake.inspect(self.source, self.out)
        summary = manifest["tables"][0]["columns"][0]["descriptive"]
        self.assertEqual(summary["median"], 1e308)
        self.assertNotIn("Infinity", (self.out / "manifest.json").read_text())

    def test_order_statistics_preserve_tiny_values_beside_extreme_outlier(self):
        self.source.write_text("measurement\n1e-308\n2e-308\n1e308\n")
        manifest, _ = intake.inspect(self.source, self.out)
        summary = manifest["tables"][0]["columns"][0]["descriptive"]
        self.assertEqual(summary["median"], 2e-308)
        self.assertEqual(summary["q1"], 1.5e-308)
        self.assertEqual(summary["q3"], 5e307)
        self.assertNotIn("Infinity", (self.out / "manifest.json").read_text())

    def test_subnormal_midpoints_and_cancellation_mean_remain_nonzero(self):
        summary = intake._column("measurement", ["5e-324", "1e-323"], [])["descriptive"]
        self.assertEqual(summary["median"], 1e-323)
        cancellation = intake._column("measurement", ["-1e308", "1e308", "1e-308"], [])["descriptive"]
        self.assertEqual(cancellation["mean"], 1e-308 / 3)

    def test_unrepresentable_spread_is_null_but_quantiles_stay_finite(self):
        summary = intake._column("measurement", ["-1.7e308", "1.7e308"], [])["descriptive"]
        self.assertIsNone(summary["sample_sd"])
        self.assertEqual(summary["median"], 0.0)
        self.assertEqual(summary["q1"], -8.5e307)
        self.assertEqual(summary["q3"], 8.5e307)
        json.dumps(summary, allow_nan=False)


if __name__ == "__main__":
    unittest.main()
