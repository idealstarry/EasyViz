#!/usr/bin/env python3
"""Extract the supplied Fig. 3f summaries, retaining XLSX numeric strings."""
import argparse
import csv
import hashlib
import json
from pathlib import Path
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import load_workbook

BASE = Path(__file__).resolve().parent
DOI = "10.1038/s41586-024-07944-6"
URL = "https://media.springernature.com/original/springer-static/esm/art%3A10.1038%2Fs41586-024-07944-6/MediaObjects/41586_2024_7944_MOESM7_ESM.xlsx"
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def numeric_strings(workbook, title):
    with zipfile.ZipFile(workbook) as archive:
        root = ET.fromstring(archive.read("xl/workbook.xml"))
        sheet = next(node for node in root.find("s:sheets", NS) if node.attrib["name"] == title)
        relation_id = sheet.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]
        relations = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        target = next(node.attrib["Target"] for node in relations if node.attrib["Id"] == relation_id)
        path = target.lstrip("/") if target.startswith("/") else "xl/" + target
        document = ET.fromstring(archive.read(path))
        return {node.attrib["r"]: node.find("s:v", NS).text
                for node in document.findall(".//s:c", NS)
                if node.attrib.get("t", "n") == "n" and node.find("s:v", NS) is not None}


def extract(workbook):
    loaded = load_workbook(workbook, data_only=True)
    sheet = loaded["3f_cytokine_matrix"]
    original = numeric_strings(workbook, sheet.title)
    genes = [sheet.cell(3, column).value for column in range(2, 67)]
    if len(set(genes)) != 65 or genes != [sheet.cell(16, column).value for column in range(2, 67)]:
        raise ValueError("Both matrix blocks must have the same 65 unique literal genes")
    regions = [sheet.cell(row, 1).value for row in range(4, 14)]
    if len(set(regions)) != 10 or regions != [sheet.cell(row, 1).value for row in range(17, 27)]:
        raise ValueError("Both matrix blocks must have the same 10 unique literal CMA bins")
    records = []
    for age, start in (("Fetal", 4), ("Paediatric", 17)):
        for offset, region in enumerate(regions):
            for column, gene in enumerate(genes, 2):
                cell = sheet.cell(start + offset, column)
                value = original[cell.coordinate]
                if float(value) != cell.value or not 0 <= float(value) <= 1:
                    raise ValueError(f"Invalid supplied normalized summary: {cell.coordinate}")
                records.append([age, region, gene, value, sheet.title, cell.coordinate])
    summary = loaded["3f_anova_cosine_sim"]
    raw = numeric_strings(workbook, summary.title)
    if [summary.cell(row, 1).value for row in range(2, 67)] != genes:
        raise ValueError("Summary genes must match the literal matrix gene order")
    summaries = []
    for row, gene in enumerate(genes, 2):
        cosine, corrected = summary.cell(row, 5), summary.cell(row, 10)
        if not -1 <= float(raw[cosine.coordinate]) <= 1 or not 0 <= float(raw[corrected.coordinate]) <= 1:
            raise ValueError("Supplied cosine / corrected interaction P is out of range")
        summaries.append([gene, raw[cosine.coordinate], raw[corrected.coordinate], summary.title,
                          cosine.coordinate, corrected.coordinate])
    return records, summaries, genes, regions


def write_csv(path, header, rows):
    with path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.writer(stream)
        writer.writerow(header)
        writer.writerows(rows)


def prepare(workbook, out):
    out.mkdir(parents=True, exist_ok=True)
    records, summaries, genes, regions = extract(workbook)
    write_csv(out / "source-data.csv", ["age", "region", "gene", "value", "source_sheet", "source_cell"], records)
    write_csv(out / "summary.csv", ["gene", "cosine", "interaction_p_adjusted", "source_sheet", "cosine_cell", "p_cell"], summaries)
    provenance = {
        "article": "A spatial human thymus cell atlas mapped to a continuous tissue axis",
        "authors": "Yayon, Kedlian, Boehme et al.", "journal": "Nature 635, 708–718 (2024)",
        "doi": DOI, "article_url": "https://www.nature.com/articles/" + DOI.split("/")[1],
        "source_url": URL, "source_workbook": workbook.name,
        "source_sha256": hashlib.sha256(workbook.read_bytes()).hexdigest(),
        "licence": "CC BY 4.0", "licence_url": "https://creativecommons.org/licenses/by/4.0/",
        "matrix_rows": len(records), "supplied_summary_rows": len(summaries),
        "numeric_source_cells": len(records) + 2 * len(summaries), "gene_order": genes, "region_order": regions,
        "transformations": ["Reshape both supplied normalized matrices into long rows; preserve original numeric XML strings and literal IDs.",
                            "Select supplied cosine similarity and corrected interaction P; no recalculation or independent-unit inference."],
        "author_code_used": False,
        "prepared_hashes": {name: hashlib.sha256((out / name).read_bytes()).hexdigest()
                            for name in ("source-data.csv", "summary.csv")},
    }
    (out / "provenance.json").write_text(json.dumps(provenance, indent=2) + "\n")
    return provenance


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=BASE / "source-fig3.xlsx")
    parser.add_argument("--out", type=Path, default=BASE)
    args = parser.parse_args()
    result = prepare(args.workbook, args.out)
    print(json.dumps({key: result[key] for key in ("matrix_rows", "supplied_summary_rows", "numeric_source_cells", "source_sha256")}))
